"""One consistent dataset for CLI exports and a future Streamlit UI."""

import csv
import json
import os
import re
import tempfile
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

from client import SpiraError, positive_id


STATUSES = {1: "Failed", 2: "Passed", 3: "Not Run", 4: "N/A", 5: "Blocked", 6: "Caution"}
COLUMNS = [
    "Project ID", "Test Set ID", "Test Set", "Folder ID", "Folder",
    "Scheduled Release ID", "Membership ID", "Position", "Test Case ID", "Test Case",
    "Configuration ID", "Configuration", "Status", "Test Run ID",
    "Execution Release ID", "Execution Release", "Completed UTC", "Build ID", "Tester ID",
    "Latest Attempt ID", "Latest Attempt Status", "Latest Attempt Started UTC",
    "Retest Pending", "Evidence",
]


def timestamp(value):
    if value in (None, ""):
        return None
    if not isinstance(value, str):
        raise SpiraError("Invalid execution timestamp.")
    legacy = re.fullmatch(r"/Date\((-?\d+)(?:[+-]\d{4})?\)/", value)
    try:
        if legacy:
            return datetime(1970, 1, 1, tzinfo=timezone.utc) + timedelta(
                milliseconds=int(legacy.group(1))
            )
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        # Spira's API contract specifies UTC even when a response omits the suffix.
        return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)
    except (ValueError, OverflowError):
        raise SpiraError(f"Unrecognized execution timestamp format: {value!r}") from None


def iso(dt):
    return dt.isoformat().replace("+00:00", "Z") if dt else ""


def status(run):
    value = run.get("ExecutionStatusId")
    if isinstance(value, bool) or value not in STATUSES:
        raise SpiraError(f"Unknown execution status ID {value!r}; update the status contract.")
    return STATUSES[value]


def configuration_id(row):
    value = row.get("TestConfigurationId")
    if value is None:
        return None
    return positive_id(row, "TestConfigurationId")


def build_report(data):
    if data.get("schema_version") != 1:
        raise SpiraError("Unsupported extraction snapshot schema.")
    project_id = positive_id(data, "project_id")
    selected_release = positive_id(data, "selected_release_id")
    releases = {positive_id(r, "ReleaseId"): r for r in data["releases"]}
    if selected_release not in releases:
        raise SpiraError("Selected release is missing from the extraction.")
    folders = {positive_id(f, "TestSetFolderId"): f for f in data["folders"]}
    sets = {positive_id(s, "TestSetId"): s for s in data["test_sets"]}
    members, configured = {}, {}
    warnings = []
    for tx, test_set in sets.items():
        for entry in data["memberships"][str(tx)]:
            mid = positive_id(entry, "TestSetTestCaseId")
            tc = positive_id(entry, "TestCaseId")
            if entry.get("TestSetId") != tx or mid in members:
                raise SpiraError("Inconsistent or duplicated membership identity.")
            case = data["test_cases"][str(tc)]
            if positive_id(case, "TestCaseId") != tc:
                raise SpiraError("Test-case response does not match membership identity.")
            members[mid] = (tx, tc, entry, case)
        configured[tx] = {}
        for entry in data["configurations"][str(tx)]:
            cid = positive_id(entry, "TestConfigurationId")
            if cid in configured[tx]:
                raise SpiraError(f"Duplicate configuration {cid} in test set {tx}.")
            configured[tx][cid] = entry
        if test_set.get("TestConfigurationSetId") and not configured[tx]:
            raise SpiraError(f"Test set {tx} has an empty configuration inventory.")
    history, run_ids = {}, set()
    excluded, unmatched = 0, 0
    for run in data["runs"]:
        rid = positive_id(run, "TestRunId")
        if rid in run_ids:
            raise SpiraError(f"Duplicate test run {rid} in extraction.")
        run_ids.add(rid)
        if run.get("ProjectId") not in (None, project_id):
            raise SpiraError("Execution response contains data from a different project.")
        if run.get("TestSetId") not in sets or run.get("ReleaseId") not in releases:
            excluded += 1
            continue
        mid = run.get("TestSetTestCaseId")
        if mid not in members:
            # Historical runs for removed memberships are not current inventory.
            unmatched += 1
            if not mid:
                raise SpiraError(
                    f"Run {rid} lacks TestSetTestCaseId; cannot safely assign it to a membership."
                )
            continue
        tx, tc, _, _ = members[mid]
        if run.get("TestSetId") != tx or run.get("TestCaseId") != tc:
            raise SpiraError(f"Run {rid} conflicts with its membership identity.")
        status(run)
        timestamp(run.get("EndDate"))
        timestamp(run.get("StartDate"))
        key = (mid, configuration_id(run))
        history.setdefault(key, []).append(run)
    if excluded:
        warnings.append(f"Excluded {excluded} returned runs outside the selected release/test-set scope.")
    if unmatched:
        warnings.append(f"Excluded {unmatched} historical runs for memberships absent from current inventory.")
    keys = set(history)
    for mid, (tx, _, _, _) in members.items():
        if configured[tx]:
            keys.update((mid, cid) for cid in configured[tx])
        elif not any(key[0] == mid for key in history):
            keys.add((mid, None))
    rows = []
    minimum = datetime.min.replace(tzinfo=timezone.utc)
    for mid, cid in sorted(keys, key=lambda k: (members[k[0]][0], k[0], k[1] or 0)):
        tx, tc, membership, case = members[mid]
        test_set = sets[tx]
        runs = history.get((mid, cid), [])
        completed = [r for r in runs if timestamp(r.get("EndDate")) and status(r) != "Not Run"]
        latest = max(completed, key=lambda r: (timestamp(r["EndDate"]), r["TestRunId"])) if completed else None
        attempt = max(runs, key=lambda r: (
            timestamp(r.get("StartDate")) or minimum, r["TestRunId"]
        )) if runs else None
        pending = bool(attempt and (
            not timestamp(attempt.get("EndDate")) or status(attempt) == "Not Run"
        ) and (not latest or
               (timestamp(attempt.get("StartDate")) or minimum) >= timestamp(latest["EndDate"])))
        if cid is not None and cid not in configured[tx]:
            warnings.append(
                f"Membership {mid}: observed configuration {cid} is not in the current configured inventory."
            )
        fid = test_set.get("TestSetFolderId")
        if fid and fid not in folders:
            raise SpiraError(f"Test set {tx} references an unknown folder {fid}.")
        execution_release = releases[latest["ReleaseId"]] if latest else {}
        config = configured[tx].get(cid, {})
        rows.append({
            "Project ID": project_id, "Test Set ID": tx, "Test Set": test_set.get("Name", ""),
            "Folder ID": fid, "Folder": folders[fid].get("Name", "") if fid else "(No folder)",
            "Scheduled Release ID": test_set.get("ReleaseId"),
            "Membership ID": mid, "Position": membership.get("Position"),
            "Test Case ID": tc, "Test Case": case.get("Name", ""),
            "Configuration ID": cid, "Configuration": config.get("Name", ""),
            "Status": status(latest) if latest else "Not Run",
            "Test Run ID": latest["TestRunId"] if latest else None,
            "Execution Release ID": latest["ReleaseId"] if latest else None,
            "Execution Release": execution_release.get("VersionNumber", ""),
            "Completed UTC": iso(timestamp(latest["EndDate"])) if latest else "",
            "Build ID": latest.get("BuildId") if latest else None,
            "Tester ID": latest.get("TesterId") if latest else None,
            "Latest Attempt ID": attempt["TestRunId"] if attempt else None,
            "Latest Attempt Status": status(attempt) if attempt else "",
            "Latest Attempt Started UTC": iso(timestamp(attempt.get("StartDate"))) if attempt else "",
            "Retest Pending": pending,
            "Evidence": ("Executed on selected release" if latest["ReleaseId"] == selected_release
                         else "Executed on child release") if latest else "No completed execution in release scope",
        })
    totals = Counter(row["Status"] for row in rows)
    summary = {"Total": len(rows), **{name: totals[name] for name in STATUSES.values()}}
    if sum(summary[name] for name in STATUSES.values()) != summary["Total"]:
        raise SpiraError("Status totals do not reconcile with report rows.")
    return {
        "metadata": {
            "project_id": project_id, "selected_release_id": selected_release,
            "release_ids": sorted(releases), "selected_folder_id": data.get("selected_folder_id"),
            "test_set_ids": sorted(sets), "extracted_utc": data.get("extracted_utc"),
            "policy": "Latest completed execution by EndDate, then TestRunId; no cross-release carry-forward",
            "identity": "TestSetTestCaseId + TestConfigurationId",
            "warnings": warnings,
            "inventory": "Current membership/configuration inventory, not historical release inventory",
        },
        "summary": summary, "rows": rows,
    }


def excel_safe(value):
    # Export untrusted artifact text as text, not spreadsheet formulae.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_report(report, output):
    path = Path(output)
    if path.suffix.lower() not in (".json", ".csv", ".xlsx"):
        raise SpiraError("Output extension must be .json, .csv, or .xlsx.")
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.resolve().parent, prefix=".spira-", suffix=path.suffix, delete=False
        ) as stream:
            temporary = Path(stream.name)
        _write_report(report, temporary)
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def _write_report(report, output):
    path = Path(output)
    if path.suffix.lower() == ".json":
        path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    elif path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows({k: excel_safe(v) for k, v in row.items()} for row in report["rows"])
    elif path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            raise SpiraError("Excel export needs openpyxl: pip install -r requirements.txt") from None
        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        for key, value in report["summary"].items():
            summary.append([key, value])
        details = wb.create_sheet("Results")
        details.append(COLUMNS)
        for row in report["rows"]:
            details.append([excel_safe(row.get(col)) for col in COLUMNS])
        details.freeze_panes = "A2"
        details.auto_filter.ref = details.dimensions
        for col in details.columns:
            details.column_dimensions[col[0].column_letter].width = 22
        metadata = wb.create_sheet("Scope")
        for key, value in report["metadata"].items():
            metadata.append([key, excel_safe(json.dumps(value) if isinstance(value, (list, dict)) else value)])
        wb.save(path)
    else:
        raise SpiraError("Output extension must be .json, .csv, or .xlsx.")
