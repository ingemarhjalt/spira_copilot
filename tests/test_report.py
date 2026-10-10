import copy
import csv
import io
import json
import os
import ssl
import sys
import tempfile
import unittest
import urllib.error
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from cli import main
from client import Settings, SpiraClient, SpiraError, folder_scope
from report import build_report, export_report, timestamp


def fixture():
    def run(rid, mid, tc, status_id, start, end, release=10, cid=None):
        return {
            "ProjectId": 1, "TestRunId": rid, "TestSetId": 20, "TestCaseId": tc,
            "TestSetTestCaseId": mid, "ExecutionStatusId": status_id,
            "ReleaseId": release, "TestConfigurationId": cid,
            "StartDate": start, "EndDate": end, "BuildId": 90, "TesterId": 7,
        }
    return {
        "schema_version": 1, "project_id": 1, "selected_release_id": 10,
        "selected_folder_id": None, "extracted_utc": "2026-10-09T12:00:00Z",
        "releases": [
            {"ReleaseId": 10, "VersionNumber": "2.7.4"},
            {"ReleaseId": 11, "VersionNumber": "2.7.4-child"},
        ],
        "folders": [],
        "test_sets": [{"TestSetId": 20, "Name": "DH8", "ReleaseId": 99}],
        "memberships": {"20": [
            {"TestSetId": 20, "TestSetTestCaseId": 101, "TestCaseId": 30, "Position": 1},
            {"TestSetId": 20, "TestSetTestCaseId": 102, "TestCaseId": 30, "Position": 2},
            {"TestSetId": 20, "TestSetTestCaseId": 103, "TestCaseId": 31, "Position": 3},
        ]},
        "test_cases": {"30": {"TestCaseId": 30, "Name": "Repeated case"},
                       "31": {"TestCaseId": 31, "Name": "Never completed"}},
        "configurations": {"20": []},
        "runs": [
            run(1000, 101, 30, 1, "2026-10-01T08:00:00Z", "2026-10-01T09:00:00Z"),
            run(999, 101, 30, 2, "2026-10-02T08:00:00Z", "2026-10-02T09:00:00Z", 11),
            run(1001, 101, 30, 3, "2026-10-03T08:00:00Z", None, 11),
            run(1002, 102, 30, 1, "2026-10-02T08:00:00Z", "2026-10-02T10:00:00Z"),
            run(1003, 101, 30, 1, "2026-10-04T08:00:00Z", "2026-10-04T10:00:00Z", 99),
        ],
    }


class ReportTests(unittest.TestCase):
    def test_completion_date_replaces_failure_without_hiding_retest(self):
        report = build_report(fixture())
        rows = report["rows"]
        self.assertEqual([r["Status"] for r in rows], ["Passed", "Failed", "Not Run"])
        self.assertEqual(rows[0]["Test Run ID"], 999)
        self.assertEqual(rows[0]["Latest Attempt ID"], 1001)
        self.assertTrue(rows[0]["Retest Pending"])
        self.assertEqual(rows[0]["Execution Release"], "2.7.4-child")
        self.assertEqual(rows[0]["Scheduled Release ID"], 99)
        self.assertEqual(rows[0]["Evidence"], "Executed on child release")
        self.assertEqual(rows[0]["Folder"], "(No folder)")
        self.assertIsNone(rows[2]["Execution Release ID"])
        self.assertEqual(report["summary"]["Total"], 3)
        self.assertIn("outside", report["metadata"]["warnings"][0])

    def test_completion_tie_uses_run_id(self):
        data = fixture()
        run = copy.deepcopy(data["runs"][1])
        run.update(TestRunId=2000, ExecutionStatusId=5)
        data["runs"].append(run)
        self.assertEqual(build_report(data)["rows"][0]["Status"], "Blocked")

    def test_empty_selected_folder_has_explicit_zero_totals(self):
        data = fixture()
        data.update(test_sets=[], memberships={}, test_cases={}, configurations={}, runs=[])
        report = build_report(data)
        self.assertEqual(report["rows"], [])
        self.assertTrue(all(value == 0 for value in report["summary"].values()))

    def test_all_statuses_in_total(self):
        data = fixture()
        data["runs"] = []
        data["memberships"]["20"] = []
        for sid in range(1, 7):
            mid = 100 + sid
            data["memberships"]["20"].append({
                "TestSetId": 20, "TestSetTestCaseId": mid, "TestCaseId": 30,
            })
            data["runs"].append({
                "TestRunId": sid, "TestSetId": 20, "TestCaseId": 30,
                "TestSetTestCaseId": mid, "ReleaseId": 10, "ExecutionStatusId": sid,
                "StartDate": "2026-10-01T08:00:00Z", "EndDate": "2026-10-01T09:00:00Z",
            })
        summary = build_report(data)["summary"]
        self.assertEqual(summary["Total"], 6)
        self.assertEqual(list(summary.values()), [6, 1, 1, 1, 1, 1, 1])

    def test_expected_and_observed_configurations_preserved(self):
        data = fixture()
        data["test_sets"][0]["TestConfigurationSetId"] = 50
        data["configurations"]["20"] = [
            {"TestConfigurationId": 61, "Name": "Silent"},
            {"TestConfigurationId": 62, "Name": "Noise"},
        ]
        data["runs"] = [data["runs"][0], data["runs"][1]]
        data["runs"][0]["TestConfigurationId"] = 61
        data["runs"][1]["TestConfigurationId"] = 63
        rows = build_report(data)["rows"]
        self.assertEqual(len(rows), 7)
        self.assertEqual([(r["Configuration ID"], r["Status"]) for r in rows[:3]],
                         [(61, "Failed"), (62, "Not Run"), (63, "Passed")])

    def test_missing_membership_fails_not_case_id_fallback(self):
        data = fixture()
        data["runs"][0]["TestSetTestCaseId"] = None
        with self.assertRaisesRegex(SpiraError, "cannot safely assign"):
            build_report(data)

    def test_removed_membership_explicitly_excluded(self):
        data = fixture()
        data["runs"][0]["TestSetTestCaseId"] = 999
        report = build_report(data)
        self.assertEqual(report["summary"]["Total"], 3)
        self.assertTrue(any("absent" in w for w in report["metadata"]["warnings"]))

    def test_invalid_run_identity_status_and_date_fail(self):
        for field, value in [
            ("TestCaseId", 999), ("ExecutionStatusId", 17), ("EndDate", "bad-date"),
            ("ProjectId", 2),
        ]:
            with self.subTest(field=field):
                data = fixture()
                data["runs"][0][field] = value
                with self.assertRaises(SpiraError):
                    build_report(data)

    def test_unknown_configuration_contract_fails(self):
        data = fixture()
        data["configurations"]["20"] = [{"EntryId": 61}]
        with self.assertRaisesRegex(SpiraError, "TestConfigurationId"):
            build_report(data)

    def test_utc_and_legacy_dates(self):
        self.assertEqual(timestamp("2026-10-01T10:00:00+02:00"),
                         timestamp("2026-10-01T08:00:00Z"))
        self.assertEqual(timestamp("2026-10-01T08:00:00"), timestamp("2026-10-01T08:00:00Z"))
        self.assertEqual(timestamp("/Date(0+0200)/").year, 1970)

    def test_offline_json_csv_and_excel_exports(self):
        data = fixture()
        data["test_cases"]["30"]["Name"] = '=HYPERLINK("unsafe")'
        report = build_report(data)
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            snapshot = folder / "snapshot.json"
            snapshot.write_text(json.dumps(data), encoding="utf-8")
            output = folder / "report.json"
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(main(["offline", "--snapshot", str(snapshot),
                                       "--output", str(output)]), 0)
            self.assertEqual(json.loads(output.read_text())["summary"], report["summary"])
            csv_path = folder / "report.csv"
            export_report(report, csv_path)
            with csv_path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(len(rows), 3)
            self.assertTrue(rows[0]["Test Case"].startswith("'="))
            from openpyxl import load_workbook
            xlsx = folder / "report.xlsx"
            export_report(report, xlsx)
            wb = load_workbook(xlsx)
            self.assertEqual(wb.sheetnames, ["Summary", "Results", "Scope"])
            self.assertEqual(wb["Results"].max_row, 4)
            self.assertNotEqual(wb["Results"]["J2"].data_type, "f")
            wb.close()

    def test_failed_export_does_not_replace_existing_report(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "report.json"
            path.write_text("original")
            with patch("report._write_report", side_effect=SpiraError("export failure")):
                with self.assertRaises(SpiraError):
                    export_report(build_report(fixture()), path)
            self.assertEqual(path.read_text(), "original")
            self.assertEqual(len(list(Path(folder).iterdir())), 1)

    def test_offline_does_not_overwrite_snapshot(self):
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            self.assertEqual(main(["offline", "--snapshot", "same.json",
                                   "--output", "same.json"]), 1)


class ClientTests(unittest.TestCase):
    def test_projects_discovery_uses_service_root_and_preserves_project_scope(self):
        client = self.client()
        urls = []
        def open_request(req, timeout):
            urls.append(req.full_url)
            return io.BytesIO(b'[{"ProjectId":1,"Name":"DH8"}]')
        with patch.object(client.opener, "open", side_effect=open_request), \
                patch("client.time.sleep"):
            self.assertEqual(client.projects()[0]["Name"], "DH8")
            client.releases()
        self.assertEqual(urls[0],
                         "https://spira.example/Services/v7_0/RestService.svc/projects")
        self.assertIn("/projects/1/releases?", urls[1])
        with self.assertRaisesRegex(SpiraError, "restricted"):
            client.request("/test-sets", project_scoped=False)

    def test_tls_verification_defaults_on(self):
        settings = Settings("https://spira.example", 1, "user", "{secret}")
        self.assertTrue(settings.verify_ssl)
        with patch("client.urllib.request.HTTPSHandler") as handler, \
                patch("client.urllib.request.build_opener"):
            with redirect_stderr(io.StringIO()) as errors:
                SpiraClient(settings)
        context = handler.call_args.kwargs["context"]
        self.assertTrue(context.check_hostname)
        self.assertEqual(context.verify_mode, ssl.CERT_REQUIRED)
        self.assertEqual(errors.getvalue(), "")

    def test_disabled_tls_verification_warns_and_ignores_ca_bundle(self):
        settings = Settings("https://spira.example", 1, "user", "{secret}",
                            ca_bundle="missing.pem", verify_ssl=False)
        with patch("client.urllib.request.HTTPSHandler") as handler, \
                patch("client.urllib.request.build_opener"):
            with redirect_stderr(io.StringIO()) as errors:
                SpiraClient(settings)
        context = handler.call_args.kwargs["context"]
        self.assertFalse(context.check_hostname)
        self.assertEqual(context.verify_mode, ssl.CERT_NONE)
        self.assertIn("verification is disabled", errors.getvalue())
        self.assertNotIn("secret", errors.getvalue())

    def test_verify_ssl_config_is_explicit_and_validated(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "spira.ini"
            prefix = "[api]\nbase_url=https://spira.example\nproject_id=1\n"
            with patch.dict(os.environ, {"SPIRA_USERNAME": "user", "SPIRA_API_KEY": "{key}"}):
                for value, expected in [("", True), ("verify_ssl=false\n", False),
                                        ("verify_ssl=true\n", True)]:
                    path.write_text(prefix + value)
                    self.assertEqual(Settings.load(str(path)).verify_ssl, expected)
                path.write_text(prefix + "verify_ssl=invalid\n")
                with self.assertRaisesRegex(SpiraError, "boolean verify_ssl"):
                    Settings.load(str(path))

    def client(self, page_size=500):
        return SpiraClient(Settings("https://spira.example", 1, "user", "{secret}",
                                    page_size=page_size))

    def test_server_capped_pagination_continues_until_empty(self):
        client = self.client()
        calls = []
        pages = [[{"TestRunId": 1}, {"TestRunId": 2}], [{"TestRunId": 3}], []]
        def array(path, params, filters):
            calls.append(params)
            return pages.pop(0)
        client.array = array
        self.assertEqual(len(client.pages("/test-runs/search", "TestRunId", filters=[])), 3)
        self.assertEqual([p["starting_row"] for p in calls], [1, 3, 4])
        self.assertEqual(calls[0]["sort_field"], "TestRunId")

    def test_pagination_repeats_abort(self):
        client = self.client()
        client.array = lambda *args: [{"TestRunId": 1}]
        with self.assertRaisesRegex(SpiraError, "Repeated"):
            client.pages("/test-runs", "TestRunId")

    def test_headers_authentication_and_read_only_post(self):
        client = self.client()
        captured = []
        def open_request(req, timeout):
            captured.append(req)
            return io.BytesIO(b"[]")
        client.opener.open = open_request
        self.assertEqual(client.request("/test-runs/search", filters=[]), [])
        req = captured[0]
        self.assertNotIn("secret", req.full_url)
        self.assertEqual(req.get_header("Api-key"), "{secret}")
        self.assertEqual(req.get_method(), "POST")
        with self.assertRaisesRegex(SpiraError, "restricted"):
            client.request("/test-sets", filters=[])
        self.assertNotIn("secret", repr(client.settings))

    def test_http_failure_aborts_and_does_not_leak_body(self):
        client = self.client()
        def fail(*args, **kwargs):
            raise urllib.error.HTTPError("secret-url", 403, "secret-message", {},
                                         io.BytesIO(b"secret-response"))
        client.opener.open = fail
        with self.assertRaises(SpiraError) as ctx:
            client.request("/test-runs")
        self.assertIn("403", str(ctx.exception))
        self.assertNotIn("secret", str(ctx.exception))

    def test_diagnostic_error_redacts_credentials_and_urls(self):
        client = self.client()
        client.diagnostic_errors = True
        message = json.dumps({
            "Message": "Invalid API key %7Bsecret%7D for user at https://host/path?api-key=secret",
            "StackTrace": "must not appear",
        }).encode()
        def fail(*args, **kwargs):
            raise urllib.error.HTTPError("url", 400, "bad", {}, io.BytesIO(message))
        with patch.object(client.opener, "open", side_effect=fail):
            with self.assertRaises(SpiraError) as ctx:
                client.request("/releases")
        detail = str(ctx.exception)
        self.assertIn("Invalid API key", detail)
        self.assertNotIn("secret", detail)
        self.assertNotIn("https://", detail)
        self.assertNotIn("must not appear", detail)

    def test_diagnostic_xml_html_and_unrecognized_bodies(self):
        client = self.client()
        bodies = [
            (b'<Fault xmlns="urn:example"><faultstring>Access denied</faultstring>'
             b'<StackTrace>private</StackTrace></Fault>', "Access denied"),
            (b"<html><title>Bad Request</title><body>private</body></html>", "Bad Request"),
            (b"arbitrary private content", "format unrecognized"),
        ]
        for body, expected in bodies:
            with self.subTest(expected=expected):
                detail = client._error_detail(io.BytesIO(body))
                self.assertIn(expected, detail)
                self.assertNotIn("private", detail)

    def test_transient_error_retries(self):
        client = self.client()
        with patch.object(client.opener, "open", side_effect=[
            urllib.error.HTTPError("url", 503, "bad", {}, io.BytesIO()),
            io.BytesIO(b"[]"),
        ]) as operation, patch("client.time.sleep"):
            self.assertEqual(client.request("/test-runs"), [])
            self.assertEqual(operation.call_count, 2)

    def test_pacing_reduces_back_to_back_request_rate_by_25_percent(self):
        client = self.client()
        clock = [0.0]
        starts = []
        def sleep(seconds):
            clock[0] += seconds
        def open_request(*args, **kwargs):
            starts.append(clock[0])
            clock[0] += 3.0
            return io.BytesIO(b"[]")
        with patch("client.time.monotonic", side_effect=lambda: clock[0]), \
                patch("client.time.sleep", side_effect=sleep), \
                patch.object(client.opener, "open", side_effect=open_request):
            for path in ["/releases", "/test-set-folders", "/test-runs"]:
                client.request(path)
        self.assertEqual(starts, [0.0, 4.0, 8.0])
        self.assertEqual(3.0 / (starts[1] - starts[0]), 0.75)

    def test_retry_backoff_counts_toward_pacing_without_shortening_it(self):
        for duration, expected_start in [(3.0, 4.0), (9.0, 12.0)]:
            with self.subTest(duration=duration):
                client = self.client()
                clock, starts = [0.0], []
                def sleep(seconds):
                    clock[0] += seconds
                def open_request(*args, **kwargs):
                    starts.append(clock[0])
                    clock[0] += duration
                    if len(starts) == 1:
                        raise urllib.error.HTTPError("url", 503, "bad", {}, io.BytesIO())
                    return io.BytesIO(b"[]")
                with patch("client.time.monotonic", side_effect=lambda: clock[0]), \
                        patch("client.time.sleep", side_effect=sleep), \
                        patch.object(client.opener, "open", side_effect=open_request):
                    client.request("/test-runs")
                self.assertEqual(starts, [0.0, expected_start])

    def test_intervening_idle_time_satisfies_cooldown(self):
        client = self.client()
        clock = [0.0]
        def open_request(*args, **kwargs):
            clock[0] += 3.0
            return io.BytesIO(b"[]")
        with patch("client.time.monotonic", side_effect=lambda: clock[0]), \
                patch("client.time.sleep") as sleep, \
                patch.object(client.opener, "open", side_effect=open_request):
            client.request("/releases")
            clock[0] += 5.0
            client.request("/test-set-folders")
            sleep.assert_not_called()

    def test_folder_scope_recursive(self):
        folders = [{"TestSetFolderId": 1, "ParentTestSetFolderId": None},
                   {"TestSetFolderId": 2, "ParentTestSetFolderId": 1},
                   {"TestSetFolderId": 3, "ParentTestSetFolderId": 2},
                   {"TestSetFolderId": 4, "ParentTestSetFolderId": None}]
        self.assertEqual(folder_scope(folders, 1), {1, 2, 3})
        with self.assertRaises(SpiraError):
            folder_scope(folders, 99)

    def test_extraction_uses_membership_mapping_and_scoped_filters(self):
        client = self.client()
        data = fixture()
        objects = {
            "/releases/10": data["releases"][0], "/releases/11": data["releases"][1],
            "/test-sets/20": data["test_sets"][0],
            "/test-cases/30": data["test_cases"]["30"],
            "/test-cases/31": data["test_cases"]["31"],
        }
        arrays = {
            "/releases/10/children": [data["releases"][1]],
            "/releases/11/children": [], "/test-set-folders": [],
            "/test-sets/20/test-case-mapping": data["memberships"]["20"],
        }
        client.object = lambda path: objects[path]
        client.array = lambda path, *args: arrays[path]
        filters_used = []
        def pages(path, id_field, params=None, filters=None):
            self.assertEqual(path, "/test-runs/search")
            filters_used.extend(filters)
            return data["runs"][:4]
        client.pages = pages
        extraction = client.extract(10, [20])
        self.assertEqual(len(extraction["memberships"]["20"]), 3)
        self.assertEqual(filters_used[0]["MultiValue"]["Values"], [10, 11])
        self.assertEqual(filters_used[1]["MultiValue"]["Values"], [20])
        self.assertEqual(build_report(extraction)["summary"]["Total"], 3)

    def test_settings_credentials_only_from_environment(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "spira.ini"
            path.write_text("[api]\nbase_url=https://spira.example\nproject_id=1\n")
            with patch.dict(os.environ, {"SPIRA_USERNAME": "user", "SPIRA_API_KEY": "{key}"}):
                settings = Settings.load(str(path))
                self.assertEqual(settings.project_id, 1)
                self.assertEqual(settings.page_size, 250)
            path.write_text("#[api]\n")
            with self.assertRaisesRegex(SpiraError, "uncommented"):
                Settings.load(str(path))

    def test_folder_extraction_includes_nested_sets_without_scheduled_release_filter(self):
        client = self.client()
        data = fixture()
        folders = [{"TestSetFolderId": 1, "ParentTestSetFolderId": None},
                   {"TestSetFolderId": 2, "ParentTestSetFolderId": 1}]
        data["test_sets"][0]["TestSetFolderId"] = 2
        objects = {
            "/releases/10": data["releases"][0], "/releases/11": data["releases"][1],
            "/test-sets/20": data["test_sets"][0],
            "/test-cases/30": data["test_cases"]["30"],
            "/test-cases/31": data["test_cases"]["31"],
        }
        arrays = {
            "/releases/10/children": [data["releases"][1]],
            "/releases/11/children": [], "/test-set-folders": folders,
            "/test-sets/20/test-case-mapping": data["memberships"]["20"],
        }
        visited = []
        def pages(path, id_field, params=None, filters=None):
            visited.append(path)
            if path == "/test-runs/search":
                return data["runs"][:4]
            self.assertIsNone(params)
            return data["test_sets"] if path.endswith("/2/test-sets") else []
        client.object = lambda path: objects[path]
        client.array = lambda path, *args: arrays[path]
        client.pages = pages
        extraction = client.extract(10, folder_id=1)
        self.assertEqual(visited[:2], [
            "/test-set-folders/1/test-sets", "/test-set-folders/2/test-sets",
        ])
        self.assertEqual(build_report(extraction)["summary"]["Total"], 3)

    def test_failed_live_retrieval_never_writes_a_report(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "report.json"
            output.write_text("existing report", encoding="utf-8")
            settings = Settings("https://spira.example", 1, "user", "{key}")
            with patch("cli.Settings.load", return_value=settings), \
                    patch("cli.SpiraClient.extract", side_effect=SpiraError("page 2 failed")), \
                    redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                self.assertEqual(main([
                    "report", "--release", "10", "--test-set", "20", "--output", str(output),
                ]), 1)
            self.assertEqual(output.read_text(), "existing report")

    def test_redirects_are_rejected(self):
        from client import NoRedirect
        with self.assertRaisesRegex(SpiraError, "redirected"):
            NoRedirect().redirect_request(None, None, 302, "", {}, "https://another.example")


if __name__ == "__main__":
    unittest.main()
